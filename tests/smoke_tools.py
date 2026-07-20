"""Smoke test integrado das ferramentas do Mangaba Operator.

Roda de verdade (navegador Chromium + servidor HTTP local + arquivos
temporários) e valida as ~40 capacidades adicionadas em jul/2026.

Uso:  PYTHONPATH=. .venv/bin/python tests/smoke_tools.py
Requisitos: macOS p/ text_to_speech/clipboard (pulados fora dele); não usa Docker.
"""

import asyncio, json, shutil, subprocess, sys, tempfile

SC = tempfile.mkdtemp(prefix="mangaba_smoke_")

def _setup_fixtures():
    with open(f"{SC}/test_page.html", "w") as fh:
        fh.write("""<!DOCTYPE html>
<html><head><title>Teste Mangaba</title></head>
<body>
<h1>Página de teste</h1>
<p>palavra repetida: teste teste teste</p>
<a href="http://localhost:8701/pagina2.html">Segunda página</a>
<a href="http://localhost:8701/upload_teste.txt">Arquivo texto</a>
<a href="http://localhost:8701/pagina2.html">Segunda página (duplicado)</a>
<input id="nome" placeholder="Nome">
<button id="dlgbtn" onclick="document.getElementById('dlg').textContent = confirm('Confirma?') ? 'aceito' : 'recusado'">Abrir diálogo</button>
<div id="dlg"></div>
</body></html>""")
    with open(f"{SC}/pagina2.html", "w") as fh:
        fh.write("<!DOCTYPE html><html><head><title>Página 2</title></head><body><h1>Segunda</h1></body></html>")
    with open(f"{SC}/upload_teste.txt", "w") as fh:
        fh.write("arquivo de upload\n")
    # imagem de teste 534x258 (mesmas dimensões usadas nos asserts)
    from PIL import Image
    Image.new("RGBA", (534, 258), (249, 117, 24, 255)).save(f"{SC}/logo_transparent.png")

_setup_fixtures()

ok, fail = [], []

async def check(label, coro, expect=None, expect_error=False):
    r = await coro
    good = bool(r.error) == expect_error and (expect is None or expect in str(r.output or r.error or ""))
    (ok if good else fail).append(label)
    print(("PASS " if good else "FAIL ") + label, "->", str(r.output or r.error or "")[:120].replace("\n", " | "))
    return r

async def main():
    from app.tool.local_server import LocalServer
    from app.tool.http_client import HttpRequest
    from app.tool.arquivos import Archive
    from app.tool.read_office import ReadOffice
    from app.tool.sqlite_query import SqliteQuery
    from app.tool.image_edit import ImageEdit
    from app.tool.clipboard import Clipboard
    from app.tool.text_to_speech import TextToSpeech
    from app.tool.text_to_pdf import TextToPdf
    from app.tool.agendar import ScheduleTask
    from app.config import config

    srv = LocalServer()
    await check("local_server start", srv.execute(action="start", port=8701, dir=SC), "http://localhost:8701")
    await check("local_server status", srv.execute(action="status"), "8701")

    http = HttpRequest()
    await check("http_request GET", http.execute(url="http://localhost:8701/test_page.html"), "HTTP 200")
    await check("http_request save_to", http.execute(url="http://localhost:8701/upload_teste.txt", save_to="baixado.txt"), "salvos em")

    # gera documentos office reais pra leitura
    import docx, openpyxl, pptx
    d = docx.Document(); d.add_paragraph("Relatório da Squid Telecom"); d.save(f"{SC}/doc.docx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Vendas"; ws.append(["cidade", "total"]); ws.append(["Rio Largo", 42]); wb.save(f"{SC}/plan.xlsx")
    pr = pptx.Presentation(); s = pr.slides.add_slide(pr.slide_layouts[0]); s.shapes.title.text = "Título do slide"; pr.save(f"{SC}/apres.pptx")
    open(f"{SC}/dados.csv", "w").write("produto,preco\nfibra,99\nlink,199\nvpn,299\n")

    ro = ReadOffice()
    await check("read_office docx", ro.execute(file_path=f"{SC}/doc.docx"), "Squid")
    await check("read_office xlsx", ro.execute(file_path=f"{SC}/plan.xlsx"), "Rio Largo")
    await check("read_office pptx", ro.execute(file_path=f"{SC}/apres.pptx"), "Título do slide")
    await check("read_office csv", ro.execute(file_path=f"{SC}/dados.csv"), "fibra")

    ar = Archive()
    await check("archive zip", ar.execute(action="zip", sources=[f"{SC}/doc.docx", f"{SC}/dados.csv"], dest=f"{SC}/pacote.zip"), "2 arquivo(s)")
    await check("archive list", ar.execute(action="list", archive_path=f"{SC}/pacote.zip"), "dados.csv")
    await check("archive unzip", ar.execute(action="unzip", archive_path=f"{SC}/pacote.zip", dest=f"{SC}/extraido"), "2 entrada(s)")

    sq = SqliteQuery()
    await check("sqlite csv_import+sql", sq.execute(sql="SELECT produto, preco FROM dados WHERE CAST(preco AS INT) > 100 ORDER BY preco", db_path=f"{SC}/t.db", csv_import=f"{SC}/dados.csv"), "vpn")

    ie = ImageEdit()
    await check("image_edit info", ie.execute(action="info", image_path=f"{SC}/logo_transparent.png"), "534x258")
    await check("image_edit resize", ie.execute(action="resize", image_path=f"{SC}/logo_transparent.png", width=200, dest=f"{SC}/logo_mini.png"), "200x")

    if shutil.which("pbcopy"):
        cb = Clipboard()
        saved = subprocess.run(["pbpaste"], capture_output=True).stdout  # preserva clipboard do usuário
        await check("clipboard copy", cb.execute(action="copy", text="mangaba-teste-123"), "17 caracteres")
        await check("clipboard paste", cb.execute(action="paste"), "mangaba-teste-123")
        subprocess.run(["pbcopy"], input=saved)
    else:
        print("SKIP clipboard (exige macOS)")

    if shutil.which("say"):
        tts = TextToSpeech()
        await check("text_to_speech", tts.execute(text="Teste do Mangaba Operator concluído."), "Áudio gerado")
    else:
        print("SKIP text_to_speech (exige macOS)")

    tp = TextToPdf()
    await check("text_to_pdf", tp.execute(text="# Relatório\n\n- item um\n- item dois\n\nParágrafo final.", title="Teste", dest=f"{SC}/rel.pdf"), "PDF gerado")

    st = ScheduleTask()
    r = await check("schedule_task add", st.execute(action="add", prompt="TAREFA-TESTE-REMOVER: nada a fazer"), "enfileirada")
    await check("schedule_task list", st.execute(action="list"), "TAREFA-TESTE-REMOVER")
    # limpa a tarefa de teste pra não rodar no cron
    for f in (config.workspace_root / "fila").glob("*.task"):
        if "TAREFA-TESTE-REMOVER" in f.read_text():
            f.unlink(); print("  (tarefa de teste removida da fila)")

    # ---- navegador: 10 novas ações servidas via http ----
    from app.tool.browser_use_tool import BrowserUseTool
    bt = BrowserUseTool()
    try:
        await check("go_to_url", bt.execute(action="go_to_url", url="http://localhost:8701/test_page.html"), "Navigated")
        await check("find_on_page", bt.execute(action="find_on_page", text="teste"), "occurrence")
        await check("list_links", bt.execute(action="list_links"), "pagina2.html")
        r = await check("list_links dedup", bt.execute(action="list_links"), "2 unique")
        await check("scroll_to_bottom", bt.execute(action="scroll_to_bottom"), "Reached the bottom")
        st = await bt.get_current_state()  # constrói o mapa de índices
        elems = json.loads(st.output)["interactive_elements"]
        print(elems)
        idx_input = idx_btn = None
        for line in elems.splitlines():
            if line.startswith("[") and "<input" in line:
                idx_input = int(line.split("[")[1].split("]")[0])
            if line.startswith("[") and "dilogo" in line.replace("á", "a").replace("Abrir diá", "Abrir dia") or "Abrir" in line:
                if "button" in line: idx_btn = int(line.split("[")[1].split("]")[0])
        print("idx_input:", idx_input, "idx_btn:", idx_btn)
        await check("fill p/ clear", bt.execute(action="input_text", index=idx_input, text="apagar"), "Input")
        await check("clear_input", bt.execute(action="clear_input", index=idx_input), "Cleared")
        r = await bt.execute(action="execute_js", script="document.getElementById('nome').value")
        assert '""' in (r.output or ""), f"campo nao limpo: {r.output}"
        print("PASS clear verificado")
        ok.append("clear verificado")

        await check("set_cookies", bt.execute(action="set_cookies", text='[{"name":"sessao","value":"abc123","url":"http://localhost:8701"}]'), "1 cookie")
        await check("get_cookies", bt.execute(action="get_cookies"), "abc123")
        await check("set_viewport", bt.execute(action="set_viewport", x=390, y=844), "390x844")
        await check("download_file", bt.execute(action="download_file", url="http://localhost:8701/upload_teste.txt", path=f"{SC}/dl_browser.txt"), "Downloaded")
        assert open(f"{SC}/dl_browser.txt").read().strip() == "arquivo de upload"
        print("PASS download conteúdo confere"); ok.append("download conteudo")

        await check("handle_dialog accept", bt.execute(action="handle_dialog", text="accept"), "auto-accept")
        st2 = await bt.get_current_state()  # reconstrói índices pós-viewport
        elems2 = json.loads(st2.output)["interactive_elements"]
        print("pós-viewport:"); print(elems2)
        for line in elems2.splitlines():
            if "button" in line:
                idx_btn = int(line.split("[")[1].split("]")[0])
        print("idx_btn agora:", idx_btn)
        await check("click diálogo", bt.execute(action="click_element", index=idx_btn), "Clicked")
        r = await bt.execute(action="execute_js", script="document.getElementById('dlg').textContent")
        assert "aceito" in (r.output or ""), f"dialogo nao aceito: {r.output}"
        print("PASS diálogo aceito de verdade"); ok.append("dialogo aceito")

        # go_forward: volta e avança
        await check("go_to_url p2", bt.execute(action="go_to_url", url="http://localhost:8701/pagina2.html"), "Navigated")
        await check("go_back", bt.execute(action="go_back"), "back")
        await check("go_forward", bt.execute(action="go_forward"), "forward")
        r = await bt.execute(action="execute_js", script="location.pathname")
        assert "pagina2" in (r.output or ""), f"forward nao voltou pra p2: {r.output}"
        print("PASS go_forward verificado"); ok.append("go_forward verificado")
    finally:
        await bt.cleanup()
        await srv.execute(action="stop", port=8701)

    shutil.rmtree(SC, ignore_errors=True)
    print(f"\n{len(ok)} PASS, {len(fail)} FAIL")
    if fail:
        print("FALHAS:", fail); sys.exit(1)

asyncio.run(main())
